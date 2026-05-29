"""Unit tests for core/extract per-format dispatch + OCR gate (issue #28).

Tests build minimal real PDF / DOCX / ODT bytes, write them through blob_store,
then call extract(sha256, mime) and assert text + paragraph_breaks.
"""
from __future__ import annotations

import hashlib
import json
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document as DocxDocument
from odf.opendocument import OpenDocumentText
from odf.text import H, P

from buscasam.core import blob_store
from buscasam.core.extract import (
    ExtractedDoc,
    OCRRequired,
    PDFEncryptionError,
    _clean_keywords,
    derive_metadata,
    extract,
    probe_encrypted,
    suggest_metadata,
)
from buscasam.settings import settings


@pytest.fixture
def blob_root(tmp_path, monkeypatch):
    root = tmp_path / "blobs"
    root.mkdir()
    monkeypatch.setattr(blob_store, "BLOB_ROOT", root)
    return root


def _persist_blob(blob_root: Path, payload: bytes) -> str:
    sha256 = hashlib.sha256(payload).hexdigest()
    final = blob_root / sha256[:2] / sha256[2:4] / sha256
    final.parent.mkdir(parents=True, exist_ok=True)
    final.write_bytes(payload)
    return sha256


def _make_docx(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_odt(paragraphs: list[str]) -> bytes:
    odt = OpenDocumentText()
    for p in paragraphs:
        odt.text.addElement(P(text=p))
    buf = BytesIO()
    odt.write(buf)
    return buf.getvalue()


_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_ODT_MIME = "application/vnd.oasis.opendocument.text"


async def test_extract_docx_returns_text_with_paragraph_breaks(blob_root):
    payload = _make_docx(["Primer párrafo.", "Segundo párrafo."])
    sha = _persist_blob(blob_root, payload)

    doc = await extract(sha, _DOCX_MIME)

    assert "Primer párrafo." in doc.text
    assert "Segundo párrafo." in doc.text
    assert len(doc.paragraph_breaks) >= 2
    assert doc.page_breaks == []
    # paragraph_breaks are ascending byte offsets into text
    assert doc.paragraph_breaks == sorted(doc.paragraph_breaks)
    assert all(0 <= b <= len(doc.text) for b in doc.paragraph_breaks)


async def test_extract_odt_returns_text_with_paragraph_breaks(blob_root):
    payload = _make_odt(["Hola mundo.", "Otra línea."])
    sha = _persist_blob(blob_root, payload)

    doc = await extract(sha, _ODT_MIME)

    assert "Hola mundo." in doc.text
    assert "Otra línea." in doc.text
    assert doc.page_breaks == []


def _make_pdf(paragraphs: list[str]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for p in paragraphs:
        pdf.multi_cell(0, 10, p)
        pdf.ln(2)
    return bytes(pdf.output())


async def test_extract_pdf_with_text_layer_returns_text(blob_root):
    paragraphs = [
        "This is a thesis abstract. " * 20,
        "Body of the work. " * 20,
        "Conclusions and references. " * 20,
    ]
    payload = _make_pdf(paragraphs)
    sha = _persist_blob(blob_root, payload)

    doc = await extract(sha, "application/pdf")

    assert "thesis abstract" in doc.text
    assert "Body of the work" in doc.text


def test_probe_encrypted_accepts_plain_pdf():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "hi")
    probe_encrypted(bytes(pdf.output()))  # should not raise


def test_probe_encrypted_rejects_encrypted_pdf():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "secret")
    pdf.set_encryption(owner_password="owner", user_password="user")
    payload = bytes(pdf.output())
    with pytest.raises(PDFEncryptionError):
        probe_encrypted(payload)


def test_probe_encrypted_does_not_raise_on_corrupted_non_encrypted_input():
    """ADR-0007 §9: parse failures unrelated to encryption fall through to async."""
    probe_encrypted(b"%PDF-1.4\nnot a real pdf body")


async def test_extract_pdf_below_threshold_raises_ocr_required(blob_root):
    # Multi-page PDF with very little text → forces avg chars/page < 100.
    from fpdf import FPDF

    pdf = FPDF()
    for _ in range(5):
        pdf.add_page()
        # No actual text on the page — simulates a scanned image.
    payload = bytes(pdf.output())
    sha = _persist_blob(blob_root, payload)

    with pytest.raises(OCRRequired) as exc:
        await extract(sha, "application/pdf")
    assert exc.value.sha256 == sha


def _doc(text: str) -> ExtractedDoc:
    # Build paragraph_breaks at each "\n\n" boundary.
    breaks: list[int] = []
    cursor = 0
    for piece in text.split("\n\n"):
        cursor += len(piece)
        breaks.append(cursor)
        cursor += 2
    return ExtractedDoc(text=text, paragraph_breaks=breaks, page_breaks=[], raw_metadata={})


def test_derive_metadata_extracts_abstract_after_resumen_heading():
    doc = _doc(
        "Tesis de grado\n\n"
        "Resumen\n\n"
        "Este trabajo investiga el ciclo de vida del bosón de Higgs.\n\n"
        "Introducción\n\n"
        "El capítulo siguiente describe la metodología."
    )
    meta = derive_metadata(doc)
    assert "bosón de Higgs" in meta.abstract


def test_derive_metadata_falls_back_to_first_paragraphs_when_no_heading():
    doc = _doc(
        "Primer párrafo del cuerpo del trabajo.\n\n"
        "Segundo párrafo del cuerpo.\n\n"
        "Tercer párrafo."
    )
    meta = derive_metadata(doc)
    assert "Primer párrafo" in meta.abstract


def test_derive_metadata_empty_doc_yields_empty_abstract():
    """ADR-0007 §9: empty extraction is not a failure."""
    doc = ExtractedDoc(text="", paragraph_breaks=[], page_breaks=[], raw_metadata={})
    meta = derive_metadata(doc)
    assert meta.abstract == ""
    assert meta.keywords == []
    assert meta.fecha is None


def test_derive_metadata_picks_recent_plausible_year_near_cover_tokens():
    doc = _doc(
        "Universidad Nacional de San Martín\n\n"
        "Tesis presentada en 2024 para optar al grado.\n\n"
        "Otra cosa de 1850 sin contexto."
    )
    from datetime import date as _date
    meta = derive_metadata(doc)
    assert meta.fecha == _date(2024, 1, 1)


def test_derive_metadata_returns_some_keywords_for_real_text():
    doc = _doc(
        "Este trabajo investiga el desarrollo del aprendizaje automático "
        "aplicado al procesamiento del lenguaje natural.\n\n"
        "El aprendizaje profundo y las redes neuronales convolucionales "
        "son técnicas centrales en el campo del aprendizaje automático."
    )
    meta = derive_metadata(doc)
    assert 0 < len(meta.keywords) <= 10
    assert all(isinstance(k, str) and k for k in meta.keywords)


def test_keyword_cleanup_filters_template_noise_and_duplicates():
    assert _clean_keywords([
        "Este trabajo",
        " aprendizaje automático ",
        "Aprendizaje automático",
        "Universidad Nacional de San Martín",
        "redes neuronales",
    ]) == ["aprendizaje automático", "redes neuronales"]


async def test_suggest_metadata_uses_llm_success(monkeypatch):
    monkeypatch.setattr(settings, "metadata_llm_enabled", True)
    monkeypatch.setattr(settings, "metadata_llm_timeout_s", 60.0)
    seen_timeout: list[float] = []
    seen_format: list[dict] = []

    def handler(req: httpx.Request) -> httpx.Response:
        seen_timeout.append(req.extensions["timeout"]["connect"])
        seen_format.append(json.loads(req.read())["format"])
        return httpx.Response(
            200,
            json={
                "response": (
                    '{"abstract": "Resumen limpio generado localmente.", '
                    '"keywords": ["grafos", "Este trabajo", "redes"]}'
                )
            },
        )

    client = httpx.AsyncClient(
        transport=httpx.MockTransport(handler), base_url="http://ollama"
    )
    try:
        meta = await suggest_metadata(
            _doc("Texto sin resumen explícito sobre grafos y redes."),
            client,
        )
    finally:
        await client.aclose()

    assert meta.abstract == "Resumen limpio generado localmente."
    assert meta.keywords == ["grafos", "redes"]
    assert seen_timeout == [60.0]
    assert seen_format[0]["required"] == ["abstract", "keywords"]
    assert seen_format[0]["additionalProperties"] is False


async def test_suggest_metadata_keeps_explicit_abstract_over_llm(monkeypatch):
    monkeypatch.setattr(settings, "metadata_llm_enabled", True)

    def handler(req: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "response": (
                    '{"abstract": "Resumen inventado por el LLM.", '
                    '"keywords": ["procesamiento de lenguaje"]}'
                )
            },
        )

    client = httpx.AsyncClient(
        transport=httpx.MockTransport(handler), base_url="http://ollama"
    )
    try:
        meta = await suggest_metadata(
            _doc(
                "Resumen\n\n"
                "Resumen determinístico extraído del documento.\n\n"
                "Introducción\n\n"
                "Cuerpo."
            ),
            client,
        )
    finally:
        await client.aclose()

    assert meta.abstract == "Resumen determinístico extraído del documento."
    assert meta.keywords == ["procesamiento de lenguaje"]


async def test_suggest_metadata_timeout_falls_back(monkeypatch):
    monkeypatch.setattr(settings, "metadata_llm_enabled", True)
    doc = _doc("Primer párrafo usado como fallback.")

    def handler(req: httpx.Request) -> httpx.Response:
        raise httpx.TimeoutException("slow local model", request=req)

    client = httpx.AsyncClient(
        transport=httpx.MockTransport(handler), base_url="http://ollama"
    )
    try:
        meta = await suggest_metadata(doc, client)
    finally:
        await client.aclose()

    assert meta == derive_metadata(doc)


async def test_suggest_metadata_invalid_output_falls_back(monkeypatch):
    monkeypatch.setattr(settings, "metadata_llm_enabled", True)
    doc = _doc("Primer párrafo usado como fallback.")

    def handler(req: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"response": '{"abstract": 123}'})

    client = httpx.AsyncClient(
        transport=httpx.MockTransport(handler), base_url="http://ollama"
    )
    try:
        meta = await suggest_metadata(doc, client)
    finally:
        await client.aclose()

    assert meta == derive_metadata(doc)


async def test_suggest_metadata_portuguese_output_falls_back(monkeypatch):
    monkeypatch.setattr(settings, "metadata_llm_enabled", True)
    doc = _doc("Texto sobre modelos de series temporales y viento.")

    def handler(req: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "response": (
                    '{"abstract": "Este documento descreve a previsão de séries temporais.", '
                    '"keywords": ["redes neurais LSTM", "vento"]}'
                )
            },
        )

    client = httpx.AsyncClient(
        transport=httpx.MockTransport(handler), base_url="http://ollama"
    )
    try:
        meta = await suggest_metadata(doc, client)
    finally:
        await client.aclose()

    assert meta == derive_metadata(doc)
